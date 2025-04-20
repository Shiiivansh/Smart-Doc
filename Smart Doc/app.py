from flask import Flask, render_template, request, jsonify, send_from_directory
import os
import requests
from fpdf import FPDF
import docx

app = Flask(__name__)
UPLOAD_FOLDER = 'uploads'
TRANSFORMED_FOLDER = os.path.join(UPLOAD_FOLDER, 'transformed')
ML_SERVICE_URL = "http://localhost:5001"

# Create folders if they don't exist
if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)
if not os.path.exists(TRANSFORMED_FOLDER):
    os.makedirs(TRANSFORMED_FOLDER)

app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER
app.config['TRANSFORMED_FOLDER'] = TRANSFORMED_FOLDER
app.config['TEMPLATES_AUTO_RELOAD'] = True

uploaded_files = {}

def get_file_content(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.pdf':
        try:
            import PyPDF2
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                text = ""
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                return text
        except Exception as e:
            print(f"PDF extraction error: {str(e)}")
            return "Unable to extract text from PDF"
    elif ext == '.docx':
        try:
            doc = docx.Document(file_path)
            return '\n'.join([para.text for para in doc.paragraphs])
        except Exception as e:
            print(f"DOCX extraction error: {str(e)}")
            return "Unable to extract text from DOCX"
    else:
        try:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception:
            return "File content could not be read as text"

def transform_content(text, template):
    try:
        response = requests.post(
            f"{ML_SERVICE_URL}/transform_document",
            json={"text": text, "template": template},
            timeout=60
        )
        if response.status_code == 200:
            transformed = response.json().get("transformed_text", "")
            return transformed.strip()
        return f"{template.upper()}\n\n{text}"
    except Exception as e:
        print(f"ML service error: {e}")
        return f"{template.upper()}\n\n{text}"

def save_as_pdf(text, output_path):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.set_auto_page_break(auto=True, margin=15)
    for line in text.split('\n'):
        pdf.multi_cell(0, 10, line)
    pdf.output(output_path)

def save_as_docx(text, output_path):
    doc = docx.Document()
    for line in text.split('\n'):
        if line.strip():
            doc.add_paragraph(line)
    doc.save(output_path)

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/upload', methods=['POST'])
def upload_file():
    if 'file' not in request.files:
        return jsonify({"error": "No file part"}), 400
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No selected file"}), 400
    file_path = os.path.join(app.config['UPLOAD_FOLDER'], file.filename)
    file.save(file_path)
    uploaded_files[file.filename] = file_path
    return jsonify({"file_name": file.filename})

@app.route('/get-suggestions', methods=['GET'])
def get_suggestions():
    file_name = request.args.get("file")
    if not file_name or file_name not in uploaded_files:
        return jsonify({"error": "File not found"}), 400
    # For simplicity, just return some templates
    return jsonify({"templates": ["Project Proposal", "Resume", "Report", "Letter", "Meeting Minutes"]})

@app.route('/apply-template', methods=['POST'])
def apply_template():
    data = request.json
    file_name = data.get("file_name")
    template = data.get("template")
    if not file_name or file_name not in uploaded_files:
        return jsonify({"error": "File not found"}), 400
    if not template:
        return jsonify({"error": "Template is required"}), 400

    file_path = uploaded_files[file_name]
    original_text = get_file_content(file_path)
    transformed_text = transform_content(original_text, template)

    # For simplicity, always save as txt
    base_name, _ = os.path.splitext(file_name)
    transformed_filename = f"transformed_{base_name}.txt"
    transformed_path = os.path.join(app.config['TRANSFORMED_FOLDER'], transformed_filename)

    with open(transformed_path, 'w', encoding='utf-8', errors='ignore') as f:
        f.write(transformed_text)

    return jsonify({
        "message": f"Applied template '{template}' to file '{file_name}'",
        "download_url": f"/download-transformed-file/{transformed_filename}"
    })

@app.route('/download-transformed-file/<filename>', methods=['GET'])
def download_transformed_file(filename):
    return send_from_directory(
        app.config['TRANSFORMED_FOLDER'],
        filename,
        as_attachment=True
    )

if __name__ == '__main__':
    app.run(debug=True)
